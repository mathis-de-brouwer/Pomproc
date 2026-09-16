#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
boma_to_pomproc.py
==================
Transform a Boma product-catalog export into the Pomproc "Products" import
sheet (the target format defined in PomProc_Data_Engine_Spec.md §4.1 and the
`Products` tab of PomProc_database_templates.xlsx).
"""

from __future__ import annotations

import argparse
import json
import logging
import re
import sys
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

logger = logging.getLogger("boma_to_pomproc")

CANONICAL_CATEGORIES = [
    "EPI", "Consommables", "Matériel", "Outillage", "Machines",
    "Carburant", "Poubelles / Sacs", "Vitrerie", "Échafaudage",
    "EPC / Protection", "Textile",
]

DEFAULT_RUBRIQUE_MAP = {
    "Chariots et mops":                "Matériel",
    "Equipement":                      "Matériel",
    "Équipement":                      "Matériel",
    "Machines et accessoires":         "Machines",
    "Nettoyage de vitres":             "Vitrerie",
    "Poubelles et sacs poubelles":     "Poubelles / Sacs",
    "Produits de nettoyage":           "Consommables",
    "Hygiène et équipement sanitaire": "Consommables",
    "Protection et sécurité":          "EPI",
    # Actual numbered source rubriques from Boma export:
    "1. EPI - Equipement de Protection Individuelle":    "EPI",
    "2. Consommables":                                  "Consommables",
    "3. Matériel":                                      "Matériel",
    "4. Outilliage":                                    "Outillage",
    "5. Machines":                                      "Machines",
    "7. Poubelles / Sacs":                              "Poubelles / Sacs",
    "8. Nettoyage Vitres":                              "Vitrerie",
    "9. Echelles / Echafaudages":                       "Échafaudage",
    "10. EPC - Equipement de Protection Collective":    "EPC / Protection",
}

HEADER_SYNONYMS = {
    "rubrique":    ["Rubrique", "Categorie", "Catégorie", "Category", "CategoryName"],
    "reference":   ["N° article", "No article", "Numero article", "Référence",
                    "Reference", "Ref", "Article", "SKU fournisseur"],
    "name":        ["Description", "Designation", "Désignation", "Name", "Libellé"],
    "unit":        ["Unité", "Unite", "Unit"],
    "package_qty": ["Quantité par emballage", "Quantite par emballage",
                    "Qté/emballage", "PackageQty", "Qte emballage", "Conditionnement"],
    "price":       ["Prix", "Price", "Prix (€)", "Prix HT", "Tarif"],
}

KNOWN_UNMAPPED_SOURCE_COLS = [
    "Certificats", "EcoCert", "EU ecolabel", "Nordic ecolabel",
    "Cradle to cradle", "Platsic Second Life", "Plastic Second Life",
    "FSC", "Approprié alimentaire", "Recycled", "Eco Score",
]

OUTPUT_COLUMNS = [
    "SKU", "Reference", "Name", "CategoryName", "Unit", "PackageQty",
    "Price", "Entity", "Active", "ImageURL", "Associations",
    "Variants", "Fournisseurs",
]

VALID_ENTITIES = {"FFA", "IIS", "DC"}

def _strip_accents(text: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFKD", text)
        if not unicodedata.combining(c)
    )

def _norm_header(text: str) -> str:
    return re.sub(r"\s+", " ", _strip_accents(str(text)).strip().lower())

def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()

def parse_decimal_eu(raw: Any) -> float | None:
    if raw is None:
        return None
    s = str(raw).strip()
    if s == "":
        return None
    s = s.replace("\u00a0", "").replace(" ", "")
    s = re.sub(r"[€$£]", "", s)

    has_comma, has_dot = "," in s, "." in s
    if has_comma and has_dot:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif has_comma:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None

def parse_int(raw: Any) -> int | None:
    val = parse_decimal_eu(raw)
    if val is None:
        return None
    return int(round(val))

def slugify(text: str, max_len: int = 28) -> str:
    s = _strip_accents(text).upper()
    s = re.sub(r"[^A-Z0-9]+", "-", s).strip("-")
    if len(s) > max_len:
        s = s[:max_len].rstrip("-")
    return s or "ITEM"

def detect_format(path: Path) -> str:
    head = path.read_bytes()[:2048]
    if head[:2] == b"PK":
        return "xlsx"
    if head[:4] == b"\xd0\xcf\x11\xe0":
        return "xls"
    lowered = head.lstrip().lower()
    if lowered.startswith(b"<") or b"<table" in lowered or b"<html" in lowered:
        return "html"
    return "html"

def _decode_bytes(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")

def read_source_rows(path: Path) -> list[dict[str, str]]:
    fmt = detect_format(path)
    logger.info("Detected source format: %s", fmt.upper())
    if fmt == "html":
        return _read_html_rows(path)
    return _read_excel_rows(path, fmt)

def _read_html_rows(path: Path) -> list[dict[str, str]]:
    from lxml import html as lxml_html
    text = _decode_bytes(path.read_bytes())
    doc = lxml_html.fromstring(text)

    matrix: list[list[str]] = []
    for tr in doc.xpath("//tr"):
        cells = [clean_text(c.text_content()) for c in tr.xpath("./th|./td")]
        if any(cells):
            matrix.append(cells)
    if not matrix:
        raise ValueError("No table rows found in the HTML export.")

    headers = matrix[0]
    rows: list[dict[str, str]] = []
    for cells in matrix[1:]:
        cells = (cells + [""] * len(headers))[:len(headers)]
        rows.append(dict(zip(headers, cells)))
    logger.info("Parsed %d data rows from HTML (headers: %d)", len(rows), len(headers))
    return rows

def _read_excel_rows(path: Path, fmt: str) -> list[dict[str, str]]:
    import pandas as pd
    engine = "openpyxl" if fmt == "xlsx" else "xlrd"
    df = pd.read_excel(path, dtype=str, engine=engine)
    df = df.where(pd.notna(df), "")
    rows = [{str(k): clean_text(v) for k, v in rec.items()}
            for rec in df.to_dict(orient="records")]
    logger.info("Parsed %d data rows from %s", len(rows), fmt.upper())
    return rows

def resolve_headers(sample_row: dict[str, str]) -> tuple[dict[str, str], list[str]]:
    lookup = {}
    for field_name, synonyms in HEADER_SYNONYMS.items():
        for syn in synonyms:
            lookup[_norm_header(syn)] = field_name

    resolved: dict[str, str] = {}
    unknown: list[str] = []
    for source_header in sample_row:
        field_name = lookup.get(_norm_header(source_header))
        if field_name and field_name not in resolved:
            resolved[field_name] = source_header
        elif field_name is None:
            unknown.append(source_header)

    missing = [f for f in ("reference", "name", "rubrique", "unit", "package_qty", "price")
               if f not in resolved]
    if missing:
        raise ValueError(
            "Source is missing required column(s): "
            + ", ".join(missing)
            + f". Headers seen: {list(sample_row)}"
        )
    logger.info("Header mapping resolved: %s", resolved)
    return resolved, unknown

@dataclass
class Reject:
    row_number: int
    reference: str
    name: str
    reason: str

@dataclass
class TransformResult:
    products: list[dict[str, Any]] = field(default_factory=list)
    rejects: list[Reject] = field(default_factory=list)
    duplicates_dropped: int = 0
    category_counts: dict[str, int] = field(default_factory=dict)

def make_sku(name: str, reference: str, used: set[str]) -> str:
    base = slugify(name)
    ref_token = re.sub(r"[^A-Z0-9]+", "", _strip_accents(reference).upper()) or "NA"
    sku = f"{base}-{ref_token}"
    candidate, n = sku, 2
    while candidate in used:
        candidate = f"{sku}-{n}"
        n += 1
    used.add(candidate)
    return candidate

def transform(
    rows: list[dict[str, str]],
    header_map: dict[str, str],
    *,
    rubrique_map: dict[str, str],
    known_categories: set[str],
    default_entity: str,
    supplier: str,
    default_category: str | None,
    strict_categories: bool,
) -> TransformResult:
    result = TransformResult()
    used_skus: set[str] = set()
    seen_rows: set[tuple] = set()

    norm_rubrique_map = {_norm_header(k): v for k, v in rubrique_map.items()}

    for i, row in enumerate(rows, start=2):
        get = lambda f: row.get(header_map.get(f, ""), "")

        reference = clean_text(get("reference"))
        name = clean_text(get("name"))
        rubrique = clean_text(get("rubrique"))
        unit = clean_text(get("unit"))
        raw_price = get("price")
        raw_pkg = get("package_qty")

        if not name:
            result.rejects.append(Reject(i, reference, name, "missing Name/Description"))
            continue
        if not reference:
            result.rejects.append(Reject(i, reference, name, "missing Reference (N° article)"))
            continue
        if not rubrique:
            result.rejects.append(Reject(i, reference, name, "missing Rubrique/Category"))
            continue

        category = norm_rubrique_map.get(_norm_header(rubrique))
        if category is None:
            if default_category:
                category = default_category
                logger.warning("Row %d: unmapped rubrique %r -> default %r",
                               i, rubrique, default_category)
            else:
                result.rejects.append(Reject(i, reference, name,
                                             f"unmapped rubrique '{rubrique}'"))
                continue
        if strict_categories and category not in known_categories:
            result.rejects.append(Reject(i, reference, name,
                                         f"category '{category}' not in Pomproc categories"))
            continue

        price = parse_decimal_eu(raw_price)
        if price is None:
            price = 0.0
            logger.warning("Row %d (%s): blank/invalid price %r -> 0.0", i, reference, raw_price)
        if price < 0:
            result.rejects.append(Reject(i, reference, name, f"negative price {price}"))
            continue

        pkg = parse_int(raw_pkg)
        if pkg is None:
            pkg = 1
        if pkg < 1:
            logger.warning("Row %d (%s): PackageQty %s < 1 -> 1", i, reference, pkg)
            pkg = 1

        if not unit:
            unit = "Pièce"
            logger.warning("Row %d (%s): blank Unit -> 'Pièce'", i, reference)

        dedup_key = (reference, name, category, unit, pkg, round(price, 4))
        if dedup_key in seen_rows:
            result.duplicates_dropped += 1
            logger.info("Row %d: exact duplicate of an earlier row (%s) — dropped",
                        i, reference)
            continue
        seen_rows.add(dedup_key)

        record = {
            "SKU":          make_sku(name, reference, used_skus),
            "Reference":    reference,
            "Name":         name,
            "CategoryName": category,
            "Unit":         unit,
            "PackageQty":   pkg,
            "Price":        round(price, 4),
            "Entity":       default_entity,
            "Active":       "yes",
            "ImageURL":     "",
            "Associations": "",
            "Variants":     "",
            "Fournisseurs": supplier,
        }
        result.products.append(record)
        result.category_counts[category] = result.category_counts.get(category, 0) + 1

    return result

def write_products_xlsx(products: list[dict[str, Any]], out_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "Products"

    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2D6CDF")
    body_font = Font(name="Arial")

    ws.append(OUTPUT_COLUMNS)
    for c in ws[1]:
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(vertical="center")

    text_cols = {"SKU", "Reference", "Associations", "Variants",
                 "ImageURL", "Fournisseurs", "Entity", "Active",
                 "CategoryName", "Unit", "Name"}

    for rec in products:
        ws.append([rec[col] for col in OUTPUT_COLUMNS])
        r = ws.max_row
        for idx, col in enumerate(OUTPUT_COLUMNS, start=1):
            cell = ws.cell(row=r, column=idx)
            cell.font = body_font
            if col in text_cols:
                cell.number_format = "@"
            elif col == "Price":
                cell.number_format = "0.####"
            elif col == "PackageQty":
                cell.number_format = "0"

    ws.freeze_panes = "A2"
    widths = {"A": 30, "B": 14, "C": 46, "D": 18, "E": 10, "F": 12,
              "G": 10, "H": 8, "I": 8, "J": 26, "K": 22, "L": 24, "M": 14}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    logger.info("Wrote %d products -> %s", len(products), out_path)

def write_rejects_xlsx(rejects: list[Reject], out_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Rejects"
    ws.append(["SourceRow", "Reference", "Name", "Reason"])
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="DC2626")
    for rj in rejects:
        ws.append([rj.row_number, rj.reference, rj.name, rj.reason])
    for col, w in {"A": 12, "B": 14, "C": 46, "D": 44}.items():
        ws.column_dimensions[col].width = w
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    logger.info("Wrote %d rejected rows -> %s", len(rejects), out_path)

def load_known_categories(categories_file: Path | None) -> set[str]:
    if not categories_file:
        return set(CANONICAL_CATEGORIES)
    try:
        from openpyxl import load_workbook
        wb = load_workbook(categories_file, data_only=True, read_only=True)
        if "Categories" not in wb.sheetnames:
            logger.warning("%s has no 'Categories' sheet; using built-in list",
                           categories_file)
            return set(CANONICAL_CATEGORIES)
        ws = wb["Categories"]
        names = set()
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                continue
            if row and row[0]:
                names.add(str(row[0]).strip())
        logger.info("Loaded %d known categories from %s", len(names), categories_file)
        return names or set(CANONICAL_CATEGORIES)
    except Exception as exc:
        logger.warning("Could not read categories from %s (%s); using built-in list",
                       categories_file, exc)
        return set(CANONICAL_CATEGORIES)

def load_rubrique_map(mapping_file: Path | None) -> dict[str, str]:
    mapping = dict(DEFAULT_RUBRIQUE_MAP)
    if mapping_file:
        user_map = json.loads(Path(mapping_file).read_text(encoding="utf-8"))
        mapping.update({str(k): str(v) for k, v in user_map.items()})
        logger.info("Applied %d rubrique overrides from %s", len(user_map), mapping_file)
    return mapping

def build_arg_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description="Transform a Boma catalog export into a Pomproc Products import sheet.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    p.add_argument("--input", "-i", required=True, type=Path,
                   help="Path to the Boma export (.xls HTML, .xlsx or legacy .xls).")
    p.add_argument("--output", "-o", type=Path, default=Path("products_pomproc.xlsx"),
                   help="Path for the Pomproc-ready Products workbook.")
    p.add_argument("--rejects", type=Path, default=Path("rejects.xlsx"),
                   help="Path for the skipped/malformed-rows report.")
    p.add_argument("--categories-file", type=Path, default=None,
                   help="Template workbook whose 'Categories' sheet lists valid categories.")
    p.add_argument("--mapping", type=Path, default=None,
                   help="JSON file of {rubrique: category} overrides.")
    p.add_argument("--default-entity", default="IIS", choices=sorted(VALID_ENTITIES),
                   help="Value written to the Entity column.")
    p.add_argument("--supplier", default="Boma",
                   help="Value written to the Fournisseurs column.")
    p.add_argument("--default-category", default=None,
                   help="Fallback category for unmapped rubriques (omit = reject them).")
    p.add_argument("--lenient-categories", action="store_true",
                   help="Do not reject rows whose mapped category is not in the known set.")
    p.add_argument("--log-level", default="INFO",
                   choices=["DEBUG", "INFO", "WARNING", "ERROR"])
    return p

def main(argv: list[str] | None = None) -> int:
    args = build_arg_parser().parse_args(argv)
    logging.basicConfig(
        level=getattr(logging, args.log_level),
        format="%(asctime)s %(levelname)-7s %(message)s",
        datefmt="%H:%M:%S",
        stream=sys.stderr,
    )

    if not args.input.exists():
        logger.error("Input file not found: %s", args.input)
        return 2

    try:
        rows = read_source_rows(args.input)
        if not rows:
            logger.error("No data rows found in the source.")
            return 1
        header_map, _unknown = resolve_headers(rows[0])
        known_categories = load_known_categories(args.categories_file)
        rubrique_map = load_rubrique_map(args.mapping)

        res = transform(
            rows, header_map,
            rubrique_map=rubrique_map,
            known_categories=known_categories,
            default_entity=args.default_entity,
            supplier=args.supplier,
            default_category=args.default_category,
            strict_categories=not args.lenient_categories,
        )

        write_products_xlsx(res.products, args.output)
        if res.rejects:
            write_rejects_xlsx(res.rejects, args.rejects)

        logger.info("=" * 60)
        logger.info("SUMMARY")
        logger.info("  source rows read      : %d", len(rows))
        logger.info("  products written      : %d", len(res.products))
        logger.info("  exact duplicates drop : %d", res.duplicates_dropped)
        logger.info("  rows rejected         : %d", len(res.rejects))
        logger.info("  categories produced   :")
        for cat, n in sorted(res.category_counts.items(), key=lambda kv: -kv[1]):
            logger.info("      %-20s %d", cat, n)
        logger.info("=" * 60)
        return 0

    except Exception as exc:
        logger.exception("Fatal error: %s", exc)
        return 1

if __name__ == "__main__":
    raise SystemExit(main())